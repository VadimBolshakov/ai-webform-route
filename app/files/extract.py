"""Pure extraction of text / images from uploaded attachments (no DB, no HTTP)."""

import asyncio
import base64
import io
import warnings
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass
from typing import Literal

from PIL import Image, ImageOps
from pypdf import PdfReader

Image.MAX_IMAGE_PIXELS = 40_000_000

TIMEOUT_S = 20.0
MAX_TEXT_CHARS = 100_000
TRUNC_MARK = "\n[…обрезано]"
MAX_IMAGE_SIDE = 2048
MAX_ZIP_UNCOMPRESSED = 50 * 1024 * 1024
IMAGE_FORMATS = {"PNG", "JPEG", "WEBP", "GIF"}


class AttachmentError(Exception):
    """Message is user-facing (Russian)."""


@dataclass
class Attachment:
    name: str
    kind: Literal["text", "image"]
    text: str | None = None
    image_data_url: str | None = None


def _truncate(text: str) -> str:
    if len(text) > MAX_TEXT_CHARS:
        return text[:MAX_TEXT_CHARS] + TRUNC_MARK
    return text


def _text_attachment(name: str, text: str) -> Attachment:
    return Attachment(name=name, kind="text", text=_truncate(text))


def _extract_pdf(name: str, data: bytes) -> Attachment:
    try:
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in reader.pages).strip()
    except Exception as exc:
        raise AttachmentError("Не удалось прочитать PDF") from exc
    if not text:
        raise AttachmentError("В PDF нет текстового слоя")
    return _text_attachment(name, text)


def _extract_docx(name: str, data: bytes) -> Attachment:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append("\t".join(c.text for c in row.cells))
        text = "\n".join(parts).strip()
    except Exception as exc:
        raise AttachmentError("Не удалось прочитать DOCX") from exc
    if not text:
        raise AttachmentError("В документе нет текста")
    return _text_attachment(name, text)


def _extract_zip(name: str, data: bytes) -> Attachment:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            is_docx = any(i.filename == "word/document.xml" for i in infos)
            total = sum(i.file_size for i in infos)
    except zipfile.BadZipFile as exc:
        raise AttachmentError("Бинарный файл не поддерживается") from exc
    if not is_docx:
        raise AttachmentError("Бинарный файл не поддерживается")
    if total > MAX_ZIP_UNCOMPRESSED:
        raise AttachmentError("Файл слишком большой после распаковки")
    return _extract_docx(name, data)


_BOMB = (Image.DecompressionBombError, Image.DecompressionBombWarning)


@contextmanager
def _bomb_as_error():
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        yield


def _sniff_image(data: bytes) -> bool:
    try:
        with _bomb_as_error(), Image.open(io.BytesIO(data)) as img:
            if img.format not in IMAGE_FORMATS:
                return False
            img.verify()
        return True
    except _BOMB as exc:
        raise AttachmentError("Изображение слишком большое") from exc
    except Exception:
        return False


def _extract_image(name: str, data: bytes) -> Attachment:
    too_big = AttachmentError("Изображение слишком большое")
    try:
        with _bomb_as_error(), Image.open(io.BytesIO(data)) as img:
            if img.width * img.height > Image.MAX_IMAGE_PIXELS:
                raise too_big
            img.seek(0)  # animated GIF/WebP: first frame only
            img = ImageOps.exif_transpose(img)  # apply orientation before EXIF is dropped
            img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
            has_alpha = img.mode in ("RGBA", "LA") or (
                img.mode == "P" and "transparency" in img.info
            )
            buf = io.BytesIO()
            if has_alpha:
                img.convert("RGBA").save(buf, "PNG")
                mime = "image/png"
            else:
                img.convert("RGB").save(buf, "JPEG", quality=85)
                mime = "image/jpeg"
    except AttachmentError:
        raise
    except _BOMB as exc:
        raise too_big from exc
    except Exception as exc:
        raise AttachmentError("Не удалось обработать изображение") from exc
    b64 = base64.b64encode(buf.getvalue()).decode("ascii")
    return Attachment(name=name, kind="image", image_data_url=f"data:{mime};base64,{b64}")


def _extract_plain(name: str, data: bytes) -> Attachment:
    if b"\x00" in data:
        raise AttachmentError("Бинарный файл не поддерживается")
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return _text_attachment(name, data.decode(enc))
        except UnicodeDecodeError:
            continue
    raise AttachmentError("Не удалось определить кодировку")


def _extract_sync(name: str, data: bytes) -> Attachment:
    if data.startswith(b"%PDF-"):
        return _extract_pdf(name, data)
    if data.startswith(b"PK\x03\x04"):
        return _extract_zip(name, data)
    if _sniff_image(data):
        return _extract_image(name, data)
    return _extract_plain(name, data)


async def extract_attachment(name: str, data: bytes, *, max_bytes: int) -> Attachment:
    if len(data) > max_bytes:
        raise AttachmentError(
            f"Файл слишком большой (максимум {max_bytes // (1024 * 1024)} МБ)"
        )
    try:
        return await asyncio.wait_for(
            asyncio.to_thread(_extract_sync, name, data), TIMEOUT_S
        )
    except asyncio.TimeoutError as exc:
        raise AttachmentError("Файл обрабатывается слишком долго") from exc
